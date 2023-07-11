import anvil.email
import anvil.secrets
import anvil.google.auth, anvil.google.drive, anvil.google.mail
from anvil.google.drive import app_files
import anvil.users
import anvil.tables as tables
import anvil.tables.query as q
from anvil.tables import app_tables
import anvil.server
import anvil.media

import xlrd
import openpyxl
import pandas as pd
from datetime import date, datetime
import requests
from bs4 import BeautifulSoup as soup

BASE_URL = 'https://www.pdga.com/player/'
PLAYER_IMG_DEFAULT_URL = 'https://cdn.pixabay.com/photo/2015/10/05/22/37/blank-profile-picture-973460_1280.png'

@anvil.server.callable
def load_spreadsheet(file) -> None:
  with anvil.media.TempFile(file) as f:
    if file.content_type == 'text/csv':
      df = pd.read_csv(f, header=None)
    else:
      df = pd.read_excel(f, header=None)
  write_records(df.to_dict('records'))

def write_records(list_of_dicts: list[dict]):
  for d in list_of_dicts:
    app_tables.dg_events.add_row(city=d[6], country=d[8], 
                                 created_ts=datetime.now(), designation=d[2], end_date=d[4].date(),
                                 fpo_champion=d[10], governing_body=d[1],
                                 id=len(app_tables.dg_events.search())+1,
                                 lmt=datetime.now(), mpo_champion=d[9], name=d[5] , start_date=d[3].date(),
                                 state=d[7], year=d[0])

@anvil.server.callable
def write_dg_event(**kwargs):
  app_tables.dg_events.add_row(id=len(app_tables.dg_events.search())+1, created_ts=datetime.now(), lmt=datetime.now(), **kwargs)

@anvil.server.callable
def get_most_recent_event() -> tuple[str, datetime]:
  return [(r['name'], r['created_ts']) for r in app_tables.dg_events.search(tables.order_by('end_date', ascending=False))][0]

@anvil.server.callable
def write_disc_golfer(pdga_id: int, first_name: str, last_name: str, division: str):
  img_url = get_player_image_url(pdga_id=pdga_id)
  app_tables.dg_players.add_row(pdga_id=pdga_id, first_name=first_name, last_name=last_name, full_name=f'{first_name} {last_name}', division=division, photo_url=img_url)


def get_player_image_url(pdga_id: int) -> str:
    url = BASE_URL + str(pdga_id)
    r = requests.get(url).content
    s = soup(r, 'html.parser')
  
    try:
        img_url = s.find(rel="gallery-player_photo").find('img').get('src')
    except:
        img_url = PLAYER_IMG_DEFAULT_URL
  
    return img_url

@anvil.server.background_task
def update_all_dg_player_photo_urls() -> None:
  for r in app_tables.dg_players.search():
    current_online_image = get_player_image_url(r['pdga_id'])
    if r['photo_url'] != current_online_image:
      print(f"Updating image for {r['full_name']}")
      r['photo_url'] = current_online_image
    else:
      print('Online image matches stored image')
